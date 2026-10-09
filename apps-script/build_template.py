"""產生「婚禮小冊・長輩意見」試算表範本（.xlsx，上傳到 Google 雲端硬碟會自動轉成 Google 試算表），
並把同一份「意見整理」版面寫進 Code.gs（setup() 會用它重建整理頁）。

用法：python3 apps-script/build_template.py [輸出路徑.xlsx]
"""
import json, re, sys
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
SUB_HEADERS = ['送出時間', '送出編號', '稱呼', '頁面版本', '意見則數', '其他想說的話', '原本合計', '照意見合計', '檢視連結', '原始資料']
OP_HEADERS = ['送出時間', '送出編號', '意見編號', '稱呼', '頁面版本', '場合代碼', '場合', '類型', '項目代碼', '項目名稱', '看法代碼', '看法', '原本金額', '建議金額', '由誰負擔', '文字', '處理狀態', '處理說明']
STATUS = ['待處理', '採納', '部分採納', '不採納', '已回覆']
D = "'意見明細'"
S = "'意見送出紀錄'"

def q(sql, empty='（還沒有資料）'):
    return f'=IFERROR(QUERY({D}!A:R,"{sql}",1),"{empty}")'

# 意見整理：(列號, [A, B, ...])，列號固定，替每張 QUERY 結果預留空間。
# 只用整欄範圍（A:R、D:D），不要用 A2:A 這種開放範圍：上傳 .xlsx 轉檔時會變成 #NAME?／#ERROR!。
SUMMARY = {
    1: ['婚禮小冊・長輩意見整理'],
    2: ['這一頁全部是公式，「意見明細」有新資料就會自動更新，請不要在這一頁打字。要標記處理結果，請到「意見明細」最後兩欄。'],
    4: ['總覽'],
    5: ['送出次數', f'=MAX(0,COUNTA({S}!B:B)-1)'],
    6: ['意見則數', f'=MAX(0,COUNTA({D}!C:C)-1)'],
    7: ['留意見的人數', f'=MAX(0,COUNTUNIQUE({D}!D:D)-1)'],
    8: ['待處理', f'=COUNTIF({D}!Q:Q,"待處理")'],
    9: ['已採納（含部分採納）', f'=COUNTIF({D}!Q:Q,"採納")+COUNTIF({D}!Q:Q,"部分採納")'],
    10: ['不採納', f'=COUNTIF({D}!Q:Q,"不採納")'],
    11: ['最近一次送出', f'=IF(COUNTA({S}!A:A)<2,"（還沒有）",TEXT(MAX({S}!A:A),"yyyy-mm-dd hh:mm"))'],
    13: ['預算：每一筆花費的看法（則數）'],
    14: [q("select J, count(C) where H = '預算' group by J pivot L label J '項目'")],
    40: ['日期：每一場的看法（則數）'],
    41: [q("select G, count(C) where H = '日期' group by G pivot L label G '場合'")],
    50: ['午宴／晚宴：比較傾向（則數）'],
    51: [q("select L, count(C) where H = '方案傾向' group by L label L '方案', count(C) '則數'")],
    58: ['依類型（則數）'],
    59: [q("select H, count(C) where H <> '' group by H order by count(C) desc label H '類型', count(C) '則數'")],
    72: ['待處理清單（最新在上）'],
    73: [q("select A, D, G, J, L, N, P where Q = '待處理' order by A desc label A '送出時間', D '稱呼', G '場合', J '項目', L '看法', N '建議金額', P '文字' format A 'yyyy-mm-dd hh:mm'", '（目前沒有待處理的意見）')],
}
TITLES = [1, 4, 13, 40, 50, 58, 72]

GUIDE = [
    ['這份試算表怎麼用'],
    [''],
    ['1. 長輩在婚禮小冊網頁按「送出給新人」，意見會自動寫進「意見送出紀錄」和「意見明細」。'],
    ['2. 「意見整理」是自動統計：哪些花費被說太多、哪些日期不方便、午宴／晚宴的傾向、待處理清單。'],
    ['3. 處理後，到「意見明細」最後兩欄填「處理狀態」與「處理說明」。Codex 也可以用 token 讀取與回寫。'],
    [''],
    ['第一次安裝（只做一次）'],
    ['A. 選單「擴充功能 → Apps Script」，刪掉預設內容，貼上 GitHub 儲存庫 apps-script/Code.gs 的全部內容，儲存。'],
    ['B. 上方函式選 setup →「執行」→ 授權（進階 → 前往）。執行記錄會顯示讀取用 token，只交給 Codex。'],
    ['C. 「部署 → 新增部署作業 → 網頁應用程式」：執行身分「我」、誰可以存取「所有人」→ 部署 → 複製網址。'],
    ['D. 把網址填進儲存庫 config.js 的 sheetEndpoint。'],
    ['E. 回到 Apps Script，函式選 testSubmit →「執行」，「意見明細」出現一列「安裝測試」就代表正常（可刪除）。'],
    [''],
    ['注意：這份試算表含有長輩的稱呼與意見，請不要公開分享；網頁只能寫入，沒有 token 讀不到內容。'],
]

def build(out):
    wb = Workbook()
    head = Font(bold=True, color='FFFFFF'); fill = PatternFill('solid', fgColor='713B43')
    summary = wb.active; summary.title = '意見整理'
    for r, cells in SUMMARY.items():
        for c, v in enumerate(cells, start=1):
            summary.cell(row=r, column=c, value=v)
    for r in TITLES: summary.cell(row=r, column=1).font = Font(bold=True, size=12 if r != 1 else 16, color='713B43')
    summary.column_dimensions['A'].width = 34
    for name, headers in (('意見明細', OP_HEADERS), ('意見送出紀錄', SUB_HEADERS)):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for cell in ws[1]: cell.font = head; cell.fill = fill; cell.alignment = Alignment(vertical='center')
        ws.freeze_panes = 'A2'
        for i, h in enumerate(headers, start=1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 36 if h in ('文字', '項目名稱', '其他想說的話', '處理說明') else 14
    dv = DataValidation(type='list', formula1='"' + ','.join(STATUS) + '"', allow_blank=True)
    wb['意見明細'].add_data_validation(dv); dv.add('Q2:Q5000')
    guide = wb.create_sheet('使用說明')
    for row in GUIDE: guide.append(row)
    for r in (1, 7): guide.cell(row=r, column=1).font = Font(bold=True, size=13, color='713B43')
    guide.column_dimensions['A'].width = 120
    wb.save(out)

def sync_code():
    code = HERE / 'Code.gs'
    rows = [[SUMMARY.get(r, [])[i] if i < len(SUMMARY.get(r, [])) else '' for i in range(2)] for r in range(1, max(SUMMARY) + 1)]
    block = ('// ---- 由 build_template.py 產生，請勿手動修改 ----\n'
             f'var SUMMARY = {json.dumps(rows, ensure_ascii=False)};\n'
             f'var SUMMARY_TITLES = {json.dumps(TITLES)};\n'
             '// ---- 產生區塊結束 ----\n')
    text = code.read_text(encoding='utf-8')
    pattern = re.compile(r'// ---- 由 build_template\.py 產生.*?// ---- 產生區塊結束 ----\n', re.S)
    text = pattern.sub(lambda _: block, text) if pattern.search(text) else text.rstrip('\n') + '\n\n' + block
    code.write_text(text, encoding='utf-8')

if __name__ == '__main__':
    out = sys.argv[1] if len(sys.argv) > 1 else str(HERE / '婚禮小冊_長輩意見_範本.xlsx')
    build(out); sync_code(); print('wrote', out)
